"""源码 Word 的学校样式、语义结构和独立源内容核验。"""
from __future__ import annotations
import hashlib
import json
from pathlib import Path
import re
import unicodedata
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn, nsmap
from lxml import etree
import pymupdf as fitz
from pdf_word_layout import figure_regions


def compact(text):
    return ''.join(unicodedata.normalize('NFKC',text).split())


def verify_source_word(path,mapping,write_report=True):
    path=Path(path)
    doc=Document(path)
    checks=[]
    def check(name,passed,detail=None):
        checks.append({'name':name,'passed':bool(passed),'detail':detail})
    check('源码原生 Word 格式',mapping.get('schema')==3 and mapping.get('mode')=='source')
    check('Word 文件哈希',hashlib.sha256(path.read_bytes()).hexdigest()==mapping['docx_sha256'])
    for source,digest in {mapping['pdf']:mapping['pdf_sha256'],**mapping['source_sha256']}.items():
        p=Path(source)
        check('源文件 '+p.name,p.is_file() and hashlib.sha256(p.read_bytes()).hexdigest()==digest)
    check('前置页与正文分节',len(doc.sections)==6)
    for i,s in enumerate(doc.sections,1):
        actual=[s.page_width.mm,s.page_height.mm,s.top_margin.mm,s.bottom_margin.mm,s.left_margin.mm,s.right_margin.mm]
        check(f'第 {i} 节 A4 与页边距',all(abs(a-b)<.05 for a,b in zip(actual,[210,297,25,25,25,20])),actual)
        check(f'第 {i} 节页眉距离',abs(s.header_distance.mm-15)<.05)
    check('允许编辑',not doc.settings.element.xpath('./w:documentProtection[not(@w:enforcement="0" or @w:enforcement="false")]'))
    check('Word 2013 排版兼容模式',doc.settings.element.xpath(
        './w:compat/w:compatSetting[@w:name="compatibilityMode"]/@w:val')==['15'])
    check('中文标点压缩',doc.settings.element.xpath(
        './w:characterSpacingControl/@w:val')==['compressPunctuation'])
    check('表格可以扩展',not doc.element.xpath('//w:trHeight[@w:hRule="exact"]'))
    check('正文没有转换标记',not any('ECJTUWORD' in p.text for p in doc.paragraphs))
    # 固定表单仅允许在封面和声明两节；学术正文必须使用流式段落。
    section_index=0
    floating_xpath=etree.XPath('.//w:framePr | .//w:txbxContent | .//wp:anchor',namespaces=nsmap)
    section_xpath=etree.XPath('./w:pPr/w:sectPr',namespaces=nsmap)
    for child in doc.element.body:
        if section_index>=2:
            check('流式段落与正文图片',not floating_xpath(child))
        if section_xpath(child):section_index+=1
    check('真实表格数量',len(doc.tables)==mapping['table_count'])
    heading_paragraphs=[p for p in doc.paragraphs if p.style.name in ('Heading 1','Heading 2','Heading 3')]
    actual_headings=[compact(p.text) for p in heading_paragraphs]
    check('标题文字、顺序和层级',actual_headings==list(map(compact,mapping['headings'])))
    with fitz.open(mapping['pdf']) as pdf:
        source_headings=[compact(row[1]) for row in pdf.get_toc()]
        check('与 LaTeX 目录独立对照',actual_headings==source_headings)
        figures=[fitz.Rect(r) for page in pdf for r in figure_regions(page, include_images=True)]
    extents=doc.element.xpath('//wp:inline/wp:extent')
    check('流程图数量',len(extents)==len(figures))
    for i,(e,r) in enumerate(zip(extents,figures),1):
        check(f'流程图 {i} 物理尺寸',abs(int(e.get('cx'))/12700-r.width)<.1 and abs(int(e.get('cy'))/12700-r.height)<.1)
    toc=[p for p in doc.paragraphs if re.fullmatch(r'toc [123]',p.style.name,re.I)]
    check('目录条目数量',len(toc)==len(heading_paragraphs))
    for i,(p,h) in enumerate(zip(toc,heading_paragraphs),1):
        check(f'目录条目 {i}',compact(p.text.rsplit('\t',1)[0])==compact(h.text) and '\t' in p.text)
    check('原生目录域',any('TOC ' in t for t in doc.element.xpath('//w:instrText/text()')))
    check('页码从正文重新编号',doc.sections[-1]._sectPr.xpath('./w:pgNumType/@w:start')==['1'])
    for i,s in enumerate(doc.sections):
        for hf in (s.header,s.even_page_header,s.footer,s.even_page_footer):
            if i<5:check('前置页隐藏页眉页码',not ''.join(hf._element.xpath('.//w:t/text() | .//w:instrText/text()')).strip())
        if i==5:
            check('奇数页页眉',compact(s.header.paragraphs[0].text)=='华东交通大学毕业设计')
            check('偶数页页眉',bool(s.even_page_header.paragraphs[0].text.strip()))
            for footer in (s.footer,s.even_page_footer):
                check('动态 PAGE 页码',any('PAGE' in t for t in footer._element.xpath('.//w:instrText/text()')))
    # 源段落清单来自 Pandoc 写入前的 AST，不能以生成的 Word 自证内容完整。
    source_paragraphs=[]
    style_names={s.style_id:s.name for s in doc.styles}
    section_index=0
    for p in doc.element.body.xpath('.//w:p'):
        if p.xpath('./w:pPr/w:sectPr'):section_index+=1;continue
        if section_index<2:continue
        if p.xpath('./w:pPr/w:framePr'):continue
        sid=p.xpath('./w:pPr/w:pStyle/@w:val')
        style=style_names[sid[0]] if sid else 'Normal'
        if style.lower().startswith('toc') or style=='Source Code':continue
        text=''.join(p.xpath('.//w:t/text()'))
        if text.strip():source_paragraphs.append(compact(text))
    expected=list(map(compact,mapping['source_content']['paragraphs']))
    mismatch=next((i for i,(a,b) in enumerate(zip(source_paragraphs,expected)) if a!=b),None)
    detail=None if mismatch is None else {'index':mismatch,'actual':source_paragraphs[mismatch],'source':expected[mismatch]}
    check('全部源段落文字与边界',source_paragraphs==expected,detail or [len(source_paragraphs),len(expected)])
    source_code=''.join(mapping['source_content'].get('code_blocks',[]))
    actual_code=''.join(p.text for p in doc.paragraphs if p.style.name=='Source Code')
    check('源代码块内容完整',compact(actual_code)==compact(source_code))
    compiled_references=sum(len(re.findall(r'\\bibitem\b',Path(name).read_text(encoding='utf-8')))
                            for name in mapping['source_sha256'] if Path(name).suffix=='.bbl')
    reference_paragraphs=[p for p in doc.paragraphs if p.style.name=='Bibliography']
    check('编译参考文献条目与编号',len(reference_paragraphs)==compiled_references and
          all(re.match(r'^\['+str(i)+r'\]',p.text) for i,p in enumerate(reference_paragraphs,1)),
          [len(reference_paragraphs),compiled_references])
    check('原生公式数量',len(doc.element.xpath('//m:oMath'))==mapping['source_content']['math_count'])
    for record in mapping['paragraphs']:
        found=doc.element.xpath('//w:bookmarkStart[@w:name="'+record['bookmark']+'"]')
        check(record['bookmark']+' 可编辑定位',len(found)==1)
        if found:
            text=''.join(found[0].getparent().xpath('.//w:t/text() | .//m:t/text()'))
            check(record['bookmark']+' 文字',compact(text)==compact(record['text']))
    # Word 会本地化字体名称并省略与继承值相同的属性，因此按样式继承求值。
    def property_value(style,group,tag,attr='val'):
        visited=set()
        while style is not None and style.style_id not in visited:
            visited.add(style.style_id)
            values=style.element.xpath(f'./w:{group}/w:{tag}/@w:{attr}')
            if values:return values[0]
            style=style.base_style
        values=doc.styles.element.xpath(f'./w:docDefaults/w:{group}Default/w:{group}/w:{tag}/@w:{attr}')
        return values[0] if values else None
    def number(style,group,tag,attr='val',default=0):
        v=property_value(style,group,tag,attr)
        return float(v) if v is not None else default
    def boolean(style,group,tag):
        while style is not None:
            values=style.element.xpath(f'./w:{group}/w:{tag}')
            if values:return values[0].get(qn('w:val'),'1') not in ('0','false','off')
            style=style.base_style
        return False
    def first_indent(style):
        while style is not None:
            for n in style.element.xpath('./w:pPr/w:ind'):
                if n.get(qn('w:hanging')) is not None:return -float(n.get(qn('w:hanging')))/20
                if n.get(qn('w:firstLine')) is not None:return float(n.get(qn('w:firstLine')))/20
            style=style.base_style
        return 0
    specifications={
        'Normal':(12,'SimSun',19.5,24),
        'Heading 1':(16,'SimHei',26,0),'Heading 2':(14,'SimHei',22.75,0),'Heading 3':(12,'SimHei',19.5,0),
        'Abstract Title':(18,'SimHei',25,0),'English Title':(18,'Times New Roman',25,0),
        'Abstract Body':(12,'SimSun',19.5,24),'English Body':(12,'SimSun',19.5,24),
        'Image Caption':(10.5,'SimHei',17.05,0),'Table Caption':(10.5,'SimHei',17.05,0),
        'Bibliography':(10.5,'SimSun',17.05,-18.5),
        'toc 1':(12,'SimHei',19.5,0),'toc 2':(12,'SimSun',19.5,0),'toc 3':(12,'SimSun',19.5,0)}
    aliases={'SimSun':{'SimSun','宋体'},'SimHei':{'SimHei','黑体'},'Times New Roman':{'Times New Roman'}}
    for name,(size,east,line,indent) in specifications.items():
        s=next((s for s in doc.styles if s.name.lower()==name.lower()),None)
        check(name+' 样式存在',s is not None)
        if s is None:continue
        actual=[number(s,'rPr','sz')/2,property_value(s,'rPr','rFonts','eastAsia'),number(s,'pPr','spacing','line')/20,first_indent(s)]
        check(name+' 字号字体行距缩进',abs(actual[0]-size)<.01 and actual[1] in aliases[east] and abs(actual[2]-line)<.03 and actual[3]==indent,actual)
        check(name+' 西文字体',property_value(s,'rPr','rFonts','ascii')=='Times New Roman')
        check(name+' 字体不受主题覆盖',not any('theme' in k.lower() for n in s.element.xpath('./w:rPr/w:rFonts') for k in n.attrib))
        check(name+' 直立字体',not boolean(s,'rPr','i'))
        check(name+' 粗体要求',boolean(s,'rPr','b')==(name in ('Heading 1','Heading 2','English Title','toc 1')))
        if name.startswith('Heading'):
            check(name+' 标题对齐',property_value(s,'pPr','jc')==('center' if name=='Heading 1' else 'left'))
            check(name+' 标题与下段同页',boolean(s,'pPr','keepNext'))
            check(name+' 按章分页',boolean(s,'pPr','pageBreakBefore')==(name=='Heading 1'))
    for i in (1,2,3):
        s=next(s for s in doc.styles if s.name.lower()==f'toc {i}')
        check(f'目录 {i} 缩进',number(s,'pPr','ind','left')==(i-1)*240 and number(s,'pPr','ind','leftChars')==0)
    for name in ('Hyperlink','FollowedHyperlink'):
        if name in doc.styles:check(name+' 链接黑色',property_value(doc.styles[name],'rPr','color')=='000000')
    parts=[doc.element,doc.styles.element,*[hf._element for s in doc.sections for hf in (s.header,s.footer,s.even_page_header,s.even_page_footer)]]
    for part in parts:
        check('文字可见且为黑色',not part.xpath('.//w14:textFill | .//w:vanish[not(@w:val="0" or @w:val="false")] | .//w:color[@w:val="FFFFFF" or @w:val="ffffff"]'))
    with ZipFile(path) as z:
        media={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
        check('封面及流程图数据完整',media==mapping['media_sha256'])
        rels=etree.fromstring(z.read('word/_rels/document.xml.rels'))
        ids={r.get('Id') for r in rels}
        check('图片关系有效',all(r in ids for r in doc.element.xpath('//a:blip/@r:embed')))
    failures=[c for c in checks if not c['passed']]
    result={'passed':not failures,'docx_sha256':mapping['docx_sha256'],'checks':checks,
            'scope':'源段落、LaTeX 目录、Word 学校样式、结构和图像；实际编辑与逐页显示另行核验。'}
    if write_report:path.with_suffix('.check.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(f'源码 Word 检查：{len(checks)-len(failures)}/{len(checks)} 通过',flush=True)
    for c in failures:print('失败：',c['name'],c['detail'],flush=True)
    return int(bool(failures))
