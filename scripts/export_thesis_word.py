"""从本项目 LaTeX 语义生成具有学校样式的原生 Word 文档。"""
from __future__ import annotations
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from zipfile import ZipFile

import pymupdf as fitz
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from pdf_word_layout import figure_regions

ROOT = Path(__file__).resolve().parents[1]
MARK = 'ECJTUWORD'
M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'


def node(tag, **attrs):
    n = OxmlElement(tag)
    for k, v in attrs.items():
        n.set(qn('w:'+k), str(v))
    return n


def strip_comments(text):
    return re.sub(r'(?<!\\)%[^\n]*', '', text)


def select_guide_branch(text, guide):
    parts = re.split(r'(\\ifecjtuguide\b|\\else\b|\\fi\b)', text)
    result, stack, active = [], [], True
    for part in parts:
        if part == r'\ifecjtuguide':
            stack.append((active, guide))
            active = active and guide
        elif part == r'\else':
            outer, selected = stack[-1]
            active = outer and not selected
        elif part == r'\fi':
            active, _ = stack.pop()
        elif active:
            result.append(part)
    if stack:
        raise ValueError('LaTeX 条件分支未闭合。')
    return ''.join(result)


def group(text, pos):
    while pos < len(text) and text[pos].isspace():
        pos += 1
    if pos >= len(text) or text[pos] != '{':
        raise ValueError('LaTeX 参数缺少大括号：'+text[pos:pos+50])
    start, depth = pos+1, 1
    pos += 1
    while pos < len(text):
        if text[pos] == '\\':
            pos += 2
            continue
        depth += (text[pos] == '{') - (text[pos] == '}')
        if not depth:
            return text[start:pos], pos+1
        pos += 1
    raise ValueError('LaTeX 参数括号未闭合。')


def replace_command(text, command, transform):
    pattern = re.compile(r'\\'+re.escape(command)+r'(?![A-Za-z])')
    result, start = [], 0
    for match in pattern.finditer(text):
        if match.start() < start:
            continue
        value, end = group(text, match.end())
        result.extend([text[start:match.start()], transform(value)])
        start = end
    return ''.join(result)+text[start:]


def expand_source(path, root, inputs, stack=()):
    path = path.resolve()
    if path in stack:
        raise ValueError('检测到循环 include：'+str(path))
    text = strip_comments(path.read_text(encoding='utf-8'))
    inputs[str(path)] = hashlib.sha256(path.read_bytes()).hexdigest()
    def include(m):
        name = Path(m.group(1))
        candidate = root / name
        if not candidate.suffix:
            candidate = candidate.with_suffix('.tex')
        if not candidate.exists():
            candidate = path.parent / name.with_suffix('.tex')
        return '\n'+expand_source(candidate, root, inputs, (*stack, path))+'\n'
    return re.sub(r'\\(?:input|include)\s*\{([^}]+)\}', include, text)


def macro_definitions(preamble):
    """Keep author-defined macros available to Pandoc's LaTeX reader."""
    result = []
    pattern = (r'\\(?:newcommand|renewcommand|providecommand)\*?\s*'
               r'(?:\{\\[A-Za-z]+\}|\\[A-Za-z]+)'
               r'(?:\s*\[[^]]*\]){0,2}\s*(?=\{)')
    for match in re.finditer(pattern, preamble):
        _, end = group(preamble, match.end())
        result.append(preamble[match.start():end])
    return '\n'.join(result)+'\n'


def expand_listings(text, root, inputs):
    def listing(match):
        options, name = match.groups()
        path = (root/name).resolve()
        code = path.read_text(encoding='utf-8')
        inputs[str(path)] = hashlib.sha256(path.read_bytes()).hexdigest()
        if options and re.search(r'\b(?:linerange|firstline|lastline)\s*=', options):
            raise ValueError('代码摘录的行范围需先另存为源文件，再用 lstinputlisting 导出。')
        return (r'\begin{lstlisting}'+('['+options+']' if options else '')+
                '\n'+code.rstrip('\n')+'\n'+r'\end{lstlisting}')
    return re.sub(r'\\lstinputlisting(?:\[([^]]*)\])?\s*\{([^}]+)\}', listing, text)


def normalize_longtables(text, labels):
    """Keep the first header, data rows and numbered caption as a native table."""
    def convert(match):
        content = match.group(1)
        columns, end = group(content, 0)
        content = content[end:]
        caption = ''
        cap = re.search(r'\\caption\s*(?=\{)', content)
        if cap:
            title, end = group(content, cap.end())
            label = re.search(r'\\label\{([^}]+)\}', content)
            number = labels.get(label.group(1), '') if label else ''
            if not number:
                raise ValueError('长表需要 label 及本次编译的编号。')
            content = content[:cap.start()]+content[end:]
            # The caption's row terminator is not a table cell.
            content = re.sub(r'^\s*(?:\\label\{[^}]+\})?\s*\\\\', '', content, count=1)
            caption = '\n\n'+MARK+'TABLECAP表'+number+'\u2003'+title+'\n\n'
        content = re.sub(r'\\label\{[^}]+\}', '', content)
        markers = list(re.finditer(r'\\end(?:firsthead|head|lastfoot|foot)\b', content))
        if markers:
            first = next((m for m in markers if m.group() == r'\endfirsthead'), None)
            first = first or next((m for m in markers if m.group() == r'\endhead'), None)
            header = content[:first.start()] if first else ''
            data = content[markers[-1].end():]
            bottom = r'\bottomrule' if r'\bottomrule' in content else ''
            content = header+'\n'+data+'\n'+bottom
        return caption+r'\begin{tabular}{'+columns+'}\n'+content+r'\end{tabular}'+'\n\n'
    return re.sub(r'\\begin\{longtable\}(?:\[[^]]*\])?(.*?)\\end\{longtable\}',
                  convert, text, flags=re.S)


def normalize_table_columns(text):
    for match in reversed(list(re.finditer(r'\\begin\{tabular\}\s*', text))):
        columns, end = group(text, match.end())
        text = text[:match.end()]+'{'+columns.replace('Y', 'l')+'}'+text[end:]
    return text


def mark_equations(text, labels):
    def equation(match):
        content = match.group(1)
        label = re.search(r'\\label\{([^}]+)\}', content)
        if label is None:
            return match.group()
        number = labels.get(label.group(1))
        if not number:
            raise ValueError('公式需要本次编译的编号。')
        content = re.sub(r'\\label\{[^}]+\}', '', content)
        return '\n\n'+r'\['+content+r'\]'+'\n\n'+MARK+'EQNUM '+number+'\n\n'
    return re.sub(r'\\begin\{equation\}(.*?)\\end\{equation\}', equation, text, flags=re.S)


def pandoc(source, source_format='latex', target='json'):
    result = subprocess.run(['pandoc', '-f', source_format, '-t', target],
        input=source, text=True, encoding='utf-8', capture_output=True, check=True)
    return result.stdout


def plain_tex(text):
    return pandoc(text, target='plain').strip()


def inlines(text):
    return [{'t':'Str', 'c':text}]


def para(text):
    return {'t':'Para', 'c':inlines(text)}


def styled(style, blocks):
    return {'t':'Div', 'c':[['', [], [['custom-style', style]]], blocks]}


def text_of(value):
    if isinstance(value, list):
        return ''.join(text_of(x) for x in value)
    if not isinstance(value, dict):
        return ''
    kind, c = value.get('t'), value.get('c')
    if kind == 'Str':
        return c
    if kind in ('Space', 'SoftBreak', 'LineBreak'):
        return ' '
    if kind in ('Code','Math'):
        return c[-1]
    if kind in ('Span','Link','Image'):
        return text_of(c[1])
    return text_of(c)


def source_manifest(ast):
    """从写入 Word 之前的源 AST 提取独立段落清单（公式单独计数）。"""
    paragraphs=[]
    math_count=0
    code_blocks=[]
    def inline(value):
        nonlocal math_count
        if isinstance(value,list):return ''.join(inline(v) for v in value)
        if not isinstance(value,dict):return ''
        kind,c=value.get('t'),value.get('c')
        if kind=='Math':math_count+=1;return ''
        if kind=='Image':return ''
        if kind=='Str':return c
        if kind in ('Space','SoftBreak','LineBreak'):return ' '
        if kind=='Code':return c[-1]
        if kind=='Quoted':
            left,right=('“','”') if c[0]['t']=='DoubleQuote' else ('‘','’')
            return left+inline(c[1])+right
        if kind in ('Span','Link'):return inline(c[1])
        return inline(c)
    def visit(value):
        if isinstance(value,list):
            for v in value:visit(v)
        elif isinstance(value,dict):
            if value.get('t') == 'CodeBlock':
                code_blocks.append(value['c'][1])
            elif value.get('t') in ('Para','Plain','Header'):
                content=value['c'][2] if value['t']=='Header' else value['c']
                text=inline(content)
                if text.strip() and not text.startswith(MARK):paragraphs.append(text)
            else:visit(value.get('c',value.get('blocks',[])))
    visit(ast['blocks'])
    result = {'paragraphs':paragraphs,'math_count':math_count}
    if code_blocks:
        result['code_blocks'] = code_blocks
    return result


def bibliography_entries(raw):
    """BibTeX 的作者可选参数可跨行；无法解析的条目应停止导出。"""
    items = list(re.finditer(r'\\bibitem\s*(?:\[[^\]]*\]\s*)?\{([^}]+)\}', raw))
    if len(items) != len(re.findall(r'\\bibitem\b', raw)):
        raise ValueError('存在无法解析的参考文献条目。')
    end = raw.index(r'\end{thebibliography}')
    entries = []
    for i, match in enumerate(items):
        content = raw[match.end():items[i+1].start() if i+1 < len(items) else end]
        content = content.replace(r'\newblock', ' ').replace(r'\allowbreak', '')
        entries.append((match.group(1), content))
    return entries


def prepare_source(source, pdf_path, work):
    root = source.parent
    inputs = {}
    all_text = expand_source(source, root, inputs)
    main = source.read_text(encoding='utf-8')
    final = bool(re.search(r'\\documentclass\[[^\]]*\bfinal\b', main))
    all_text = select_guide_branch(all_text, not final)
    setup = re.search(r'\\thesissetup\s*', all_text)
    if not setup:
        raise ValueError('本导出入口需要 ecjtu-thesis 的 thesissetup 元数据。')
    setup_text, _ = group(all_text, setup.end())
    meta = {}
    for match in re.finditer(r'([\w-]+)\s*=\s*(?=\{)', setup_text):
        value, _ = group(setup_text, match.end())
        meta[match.group(1)] = plain_tex(value.replace(r'\quad', '\u2003'))
    preamble, text = all_text.split(r'\begin{document}',1)
    text = text.split(r'\end{document}',1)[0]
    definitions_from_source = macro_definitions(preamble)
    for command, marker in [('makecover','COVER'),('makedeclaration','DECLARATION'),
                            ('thesistableofcontents','TOC'),('thesismainmatter','BODY'),
                            ('thesisappendices','APPENDICES')]:
        text = re.sub(r'\\'+command+r'\b', '\n\n'+MARK+marker+'\n\n', text)
    for env, begin, end in [('cnabstract','CN','ENDCN'),('enabstract','EN','ENDEN')]:
        text = text.replace(r'\begin{'+env+'}', '\n\n'+MARK+begin+'\n\n')
        text = text.replace(r'\end{'+env+'}', '\n\n'+MARK+end+'\n\n')
    text = replace_command(text, 'keywords', lambda v:'\n\n'+MARK+'KEYCN '+v+'\n\n')
    text = replace_command(text, 'keywordsen', lambda v:'\n\n'+MARK+'KEYEN '+v+'\n\n')
    text = replace_command(text, 'thesisbackchapter', lambda v:r'\chapter*{'+v+'}')
    text = re.sub(r'\\(?:clearpage|newpage)\b', '\n\n'+MARK+'PAGEBREAK\n\n', text)
    for command in ('pagenumbering','pagestyle','thispagestyle','nocite'):
        text = replace_command(text, command, lambda v:'')
    # 引用显示值沿用本次 LaTeX 编译结果。
    labels, citations = {}, {}
    for path in (root/'build').rglob('*.aux'):
        raw = path.read_text(encoding='utf-8')
        inputs[str(path)] = hashlib.sha256(path.read_bytes()).hexdigest()
        labels.update(re.findall(r'\\newlabel\{([^}]+)\}\{\{([^}]+)\}',raw))
        citations.update(re.findall(r'\\bibcite\{([^}]+)\}\{\{([^}]+)\}',raw))
    text = replace_command(text, 'ref', lambda key:labels[key])
    text = replace_command(text, 'eqref', lambda key:'('+labels[key]+')')
    text = mark_equations(text, labels)
    def cite(keys):
        nums=sorted(set(int(citations[k.strip()]) for k in keys.split(',')))
        groups=[]
        for n in nums:
            if groups and n==groups[-1][-1]+1:groups[-1].append(n)
            else:groups.append([n])
        label=','.join(str(g[0]) if len(g)==1 else f'{g[0]}-{g[-1]}' for g in groups)
        return r'\textsuperscript{['+label+']}'
    text = replace_command(text, 'cite', cite)
    bibliography = []
    def bib(_):
        bbl = root/'build'/f'{source.stem}.bbl'
        inputs[str(bbl)] = hashlib.sha256(bbl.read_bytes()).hexdigest()
        raw = bbl.read_text(encoding='utf-8')
        bibliography.extend(bibliography_entries(raw))
        return '\n\n'+MARK+'BIBLIOGRAPHY\n\n'
    text = replace_command(text,'bibliography',bib)
    # 当前 PDF 中连通矢量图与源 TikZ 环境按文档顺序对应；正文仍取自源码。
    drawings=[]
    with fitz.open(pdf_path) as pdf:
        for page in pdf:
            for rect in figure_regions(page):
                drawings.append((page.number,rect))
        tikz=list(re.finditer(r'\\begin\{tikzpicture\}.*?\\end\{tikzpicture\}',text,re.S))
        if len(tikz)!=len(drawings):
            raise ValueError(f'TikZ 图与 PDF 图形数不一致：{len(tikz)} / {len(drawings)}。')
        for i, match in reversed(list(enumerate(tikz))):
            page,rect=drawings[i]
            image=work/f'figure-{i+1}.png'
            pdf[page].get_pixmap(clip=fitz.Rect(rect),dpi=300).save(image)
            width=fitz.Rect(rect).width
            text=text[:match.start()]+r'\includegraphics[width='+str(width)+'pt]{'+image.as_posix()+'}'+text[match.end():]
    text=re.sub(r'\\resizebox\{[^{}]*\}\{[^{}]*\}\s*\{\s*'
                r'(\\includegraphics(?:\[[^]]*\])?\{[^{}]+\})\s*\}', r'\1', text)
    # 每个浮动体显式保留编号、图题/表题的前后位置。
    figure_count=table_count=0
    def floating(match):
        nonlocal figure_count,table_count
        kind,content=match.group(1),match.group(2)
        cap=re.search(r'\\caption(?![A-Za-z])\s*',content)
        if not cap:
            raise ValueError('浮动体缺少 caption，无法恢复题注。')
        caption,end=group(content,cap.end())
        label=re.search(r'\\label\{([^}]+)\}',content)
        number=labels.get(label.group(1),'') if label else ''
        if not number:
            raise ValueError('图表需要 label 及本次编译的编号。')
        content=content[:cap.start()]+content[end:]
        content=re.sub(r'\\label\{[^}]+\}', '', content)
        caption='\n\n'+MARK+('FIGCAP' if kind=='figure' else 'TABLECAP')+('图' if kind=='figure' else '表')+number+'\u2003'+caption+'\n\n'
        if kind=='figure':
            figure_count+=1
            return '\n\n'+content+caption
        table_count+=1
        return caption+content+'\n\n'
    text=re.sub(r'\\begin\{(figure|table)\}(?:\[[^\]]*\])?(.*?)\\end\{\1\}',floating,text,flags=re.S)
    # Pandoc 能读取 tabular 的语义单元格；列宽稍后根据模板明确设置。
    text=re.sub(r'\\begin\{tabularx\}\{\\textwidth\}',r'\\begin{tabular}',text)
    text=text.replace(r'\end{tabularx}',r'\end{tabular}')
    text=re.sub(r'>\{\\raggedright\\arraybackslash\}', '', text)
    text=text.replace('|Y|Y|','|l|l|')
    text=text.replace(r'\tablefont','')
    text=normalize_longtables(text, labels)
    text=normalize_table_columns(text)
    text=expand_listings(text, root, inputs)
    definitions=(r'\newcommand{\templatehint}[1]{'+('' if final else r'#1\par')+'}\n'
                 r'\newcommand{\templateheading}[1]{'+('' if final else r'\paragraph*{#1}')+'}\n'
                 r'\newcommand{\templateblank}{\templatehint{（在此撰写本节内容。）}}'+'\n')
    ast=json.loads(pandoc(definitions_from_source+definitions+text))
    table_count=sum(block['t']=='Table' for block in ast['blocks'])
    blocks=[]
    mode='front'; chapter=section=subsection=0; appendix=0; headings=[]
    for block in ast['blocks']:
        string=text_of(block)
        if block['t']=='Para' and string.startswith(MARK):
            code=string[len(MARK):]
            if code=='CN':
                mode='cn'
                blocks += [para(MARK+'CN'),styled('Abstract Title',[para(meta['title'])]),styled('Front Heading',[para('摘　要')])]
            elif code=='EN':
                mode='en'
                blocks += [para(MARK+'EN'),styled('English Title',[para(meta['title-en'])]),styled('Front Heading English',[para('ABSTRACT')])]
            elif code in ('ENDCN','ENDEN'):
                mode='front'
            elif code=='BODY':
                mode='body';blocks.append(block)
            elif code=='APPENDICES':
                appendix=1
            elif code=='BIBLIOGRAPHY':
                blocks.append({'t':'Header','c':[1,['references',['unnumbered'],[]],inlines('参考文献')]})
                headings.append('参考文献')
                for i,(_,content) in enumerate(bibliography):
                    parsed=json.loads(pandoc(f'[{i+1}] '+content))['blocks']
                    merged=[]
                    for p in parsed:
                        if merged:merged.append({'t':'Space'})
                        merged.extend(p.get('c',[]))
                    blocks.append(styled('Bibliography',[{'t':'Para','c':merged}]))
            elif code.startswith(('KEYCN','KEYEN')):
                kind=code[:5]
                rest=block['c'][1:] if block['c'][0].get('c')==MARK+kind else None
                # 标记与关键词可能被 Pandoc 合并为同一个 Str。
                values=deepcopy(block['c'])
                values[0]['c']=values[0]['c'].replace(MARK+kind,('关键词：' if kind=='KEYCN' else 'Keywords:'),1)
                blocks.append(styled('Keywords' if kind=='KEYCN' else 'Keywords English',[{'t':'Para','c':values}]))
            elif code.startswith(('FIGCAP','TABLECAP')):
                values=deepcopy(block['c'])
                tag='FIGCAP' if code.startswith('FIGCAP') else 'TABLECAP'
                values[0]['c']=values[0]['c'].replace(MARK+tag,'',1)
                blocks.append(styled('Image Caption' if tag=='FIGCAP' else 'Table Caption',[{'t':'Para','c':values}]))
            elif code.startswith('EQNUM '):
                previous=blocks.pop()
                if previous['t']!='Para' or len(previous['c'])!=1 or previous['c'][0]['t']!='Math':
                    raise ValueError('无法定位编号公式的原生数学内容。')
                math=deepcopy(previous['c'][0])
                math['c'][0]={'t':'InlineMath'}
                blocks.append(styled('Equation',[{'t':'Para','c':[math,*inlines('('+code[6:].strip()+')')]}]))
            else:
                blocks.append(block)
            continue
        if block['t']=='Header':
            level,attrs,values=block['c']
            if level>=4:
                blocks.append(styled('Hint Heading',[{'t':'Para','c':values}]))
                continue
            if level==1 and 'unnumbered' in attrs[1]:
                title=text_of(values)
            elif level==1:
                section=subsection=0
                if appendix:
                    prefix='附录'+chr(64+appendix);appendix+=1
                else:
                    chapter+=1;prefix=f'第{chapter}章'
                values=inlines(prefix+'　')+values
                title=text_of(values)
            elif level==2:
                section+=1;subsection=0
                values=inlines(f'{chapter}.{section}\u2002')+values;title=text_of(values)
            else:
                subsection+=1
                values=inlines(f'{chapter}.{section}.{subsection}\u2002')+values;title=text_of(values)
            block['c'][2]=values
            headings.append(title)
        if mode in ('cn','en') and block['t']=='Para':
            block=styled('Abstract Body' if mode=='cn' else 'English Body',[block])
        blocks.append(block)
    ast['blocks']=blocks
    raw=[b for b in blocks if b['t']=='RawBlock']
    if raw:
        raise ValueError('存在未转换的 LaTeX 块：'+str(raw)[:200])
    prepared=work/'prepared.json'
    prepared.write_text(json.dumps(ast,ensure_ascii=False),encoding='utf-8')
    return prepared,meta,inputs,headings,table_count


def font(style, east='SimSun', size=12, bold=False):
    style.font.name='Times New Roman'
    style.font.size=Pt(size)
    style.font.bold=bold
    style.font.italic=False
    style.font.color.rgb=RGBColor(0,0,0)
    pr=style.element.get_or_add_rPr()
    fonts=pr.get_or_add_rFonts()
    for attr in list(fonts.attrib):
        if 'theme' in attr.lower():del fonts.attrib[attr]
    fonts.set(qn('w:eastAsia'),east)
    fonts.set(qn('w:cs'),'Times New Roman')
    old=pr.find(qn('w:szCs'))
    if old is not None:pr.remove(old)
    pr.append(node('w:szCs',val=round(size*2)))


def set_style(doc,name,size=12,east='SimSun',bold=False,line=19.5,indent=24,
              before=0,after=0,align=WD_ALIGN_PARAGRAPH.JUSTIFY,keep=False,page=False):
    style=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    if name!='Normal':style.base_style=doc.styles['Normal']
    style.next_paragraph_style=doc.styles['Normal']
    font(style,east,size,bold)
    p=style.paragraph_format
    p.first_line_indent=Pt(indent);p.left_indent=p.right_indent=Pt(0)
    p.space_before=Pt(before);p.space_after=Pt(after);p.line_spacing=Pt(line)
    p.alignment=align;p.keep_with_next=keep;p.keep_together=False;p.page_break_before=page;p.widow_control=True
    for name in ('snapToGrid','autoSpaceDE','autoSpaceDN'):
        pr=style.element.get_or_add_pPr()
        old=pr.find(qn('w:'+name))
        if old is not None:pr.remove(old)
        # xeCJK inserts a quarter-em gap at Chinese/Latin and Chinese/digit
        # boundaries. Word's East Asian auto-spacing provides that same gap.
        pr.append(node('w:'+name,val=0 if name=='snapToGrid' else 1))
    return style


def styles(doc):
    set_style(doc,'Normal')
    for name in ('Body Text','First Paragraph','Abstract Body','English Body'):
        set_style(doc,name)
    set_style(doc,'Hint Heading',keep=True)
    for level,size,line,before,after in [(1,16,26,0,19.5),(2,14,22.75,9.75,0),(3,12,19.5,0,0)]:
        set_style(doc,f'Heading {level}',size,'SimHei',level<3,line,0,before,after,
                  WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT,True,level==1)
    for name,size,east,bold,line,after in [('Abstract Title',18,'SimHei',False,25,15.6),
        ('English Title',18,'Times New Roman',True,25,15.6),('Front Heading',16,'SimHei',False,26,7.8),
        ('Front Heading English',16,'Times New Roman',True,26,7.8)]:
        set_style(doc,name,size,east,bold,line,0,0,after,WD_ALIGN_PARAGRAPH.CENTER,True)
    set_style(doc,'TOC Heading',16,'SimHei',True,26,0,0,15.6,WD_ALIGN_PARAGRAPH.CENTER,True)
    set_style(doc,'Keywords',12,'SimHei',False,19.5,0,19.5,0,WD_ALIGN_PARAGRAPH.LEFT)
    set_style(doc,'Keywords English',12,'SimSun',True,19.5,0,19.5,0,WD_ALIGN_PARAGRAPH.LEFT)
    set_style(doc,'Source Code',10.56,line=15.6,indent=0,align=WD_ALIGN_PARAGRAPH.LEFT)
    set_style(doc,'Equation',12,line=19.5,indent=0,before=9.75,after=9.75,align=WD_ALIGN_PARAGRAPH.LEFT)
    for name in ('Image Caption','Table Caption'):
        set_style(doc,name,10.5,'SimHei',False,17.0625,0,6,6,WD_ALIGN_PARAGRAPH.CENTER,name=='Table Caption')
    s=set_style(doc,'Bibliography',10.5,line=17.0625,indent=-18.5,align=WD_ALIGN_PARAGRAPH.LEFT)
    s.paragraph_format.left_indent=Pt(18.5)
    for name in ('Hyperlink','FollowedHyperlink'):
        if name in doc.styles:
            doc.styles[name].font.color.rgb=RGBColor(0,0,0)
            doc.styles[name].font.underline=False
    for level in (1,2,3):
        s=set_style(doc,f'toc {level}',12,'SimHei' if level==1 else 'SimSun',level==1,19.5,0,
                    5 if level==1 else 0,0,WD_ALIGN_PARAGRAPH.LEFT)
        s.element.set(qn('w:customStyle'),'0')
        s.paragraph_format.left_indent=Pt((level-1)*12)
    for p in doc.paragraphs:
        if p.style.name in ('Normal','Body Text','First Paragraph','Compact'):
            p.style=doc.styles['Normal']
        # Pandoc 的段落样式继承关系已替换；清理与模板相抵的直接缩进。
        for tag in ('w:ind','w:spacing','w:jc'):
            pr=p._p.pPr
            item=pr.find(qn(tag)) if pr is not None else None
            if item is not None:pr.remove(item)
        for r in p.runs:
            if r._r.rPr is not None:
                for n in list(r._r.rPr):
                    if n.tag in [qn(t) for t in ('w:rFonts','w:sz','w:szCs','w:color')]:
                        r._r.rPr.remove(n)


def configure_section(section):
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=section.bottom_margin=section.left_margin=Mm(25)
    section.right_margin=Mm(20)
    section.header_distance=Mm(15);section.footer_distance=Mm(16.75)
    for tag in ('w:docGrid','w:pgNumType','w:titlePg'):
        n=section._sectPr.find(qn(tag))
        if n is not None:section._sectPr.remove(n)


def break_before(marker, template):
    p=marker.insert_paragraph_before()
    p.paragraph_format.line_spacing=Pt(1)
    p.paragraph_format.space_after=p.paragraph_format.space_before=Pt(0)
    p.add_run().font.size=Pt(1)
    sect=deepcopy(template)
    sect.append(node('w:type',val='nextPage')) if sect.find(qn('w:type')) is None else None
    p._p.get_or_add_pPr().append(sect)


def set_run(run,size=12,east='SimSun',bold=False):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.font.bold=bold
    run.font.italic=False
    run.font.color.rgb=RGBColor(0,0,0)
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),east)


def add_cover(marker, pdf):
    page=pdf[0]
    info=page.get_image_info()[0]
    rect=fitz.Rect(info['bbox'])
    image=page.parent.extract_image(page.get_images()[0][0])['image']
    from io import BytesIO
    holder=marker.insert_paragraph_before()
    holder.paragraph_format.line_spacing=Pt(1)
    holder.paragraph_format.space_after=Pt(0)
    inline=holder.add_run().add_picture(BytesIO(image),width=Pt(rect.width),height=Pt(rect.height))._inline
    inline.tag=qn('wp:anchor')
    for k,v in dict(distT=0,distB=0,distL=0,distR=0,simplePos=0,relativeHeight=0,behindDoc=1,locked=0,layoutInCell=1,allowOverlap=1).items():inline.set(k,str(v))
    simple=OxmlElement('wp:simplePos');simple.set('x','0');simple.set('y','0');inline.insert(0,simple)
    for i,(axis,pos) in enumerate([('H',rect.x0),('V',rect.y0)],1):
        n=OxmlElement('wp:position'+axis);n.set('relativeFrom','page')
        offset=OxmlElement('wp:posOffset');offset.text=str(round(pos*12700));n.append(offset);inline.insert(i,n)
    extent=inline.find(qn('wp:extent'));inline.insert(list(inline).index(extent)+1,OxmlElement('wp:wrapNone'))
    # 封面填写项是可编辑文字；坐标取自已编译的封面，不包含正文排版推断。
    for b in page.get_text('dict')['blocks']:
        for line in b.get('lines',[]):
            p=marker.insert_paragraph_before()
            p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(0)
            size=max(s['size'] for s in line['spans'])
            p.paragraph_format.line_spacing=Pt(size*1.25)
            x0,y0,x1,y1=line['bbox']
            p._p.get_or_add_pPr().append(node('w:framePr',w=round((x1-x0+2)*20),
                h=1,hRule='atLeast',hAnchor='page',vAnchor='page',x=round(x0*20),y=round((y0-1)*20),wrap='none'))
            for s in line['spans']:
                r=p.add_run(s['text']);set_run(r,s['size'])
    marker._p.getparent().remove(marker._p)


def add_declaration(marker,doc,meta):
    from lxml import etree
    holder=marker.insert_paragraph_before()
    holder.paragraph_format.line_spacing=Pt(1)
    holder.paragraph_format.space_after=Pt(0)
    pict=OxmlElement('w:pict')
    rect=etree.SubElement(pict,'{urn:schemas-microsoft-com:vml}rect')
    rect.set('style','position:absolute;margin-left:119.055pt;margin-top:70.866pt;'
        'width:374.173pt;height:629.291pt;z-index:-251654144;'
        'mso-position-horizontal-relative:page;mso-position-vertical-relative:page')
    rect.set('filled','f');rect.set('strokecolor','#808080');rect.set('strokeweight','0.35pt')
    holder.add_run()._r.append(pict)
    def position(p,x,y,width):
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(0)
        p._p.get_or_add_pPr().append(node('w:framePr',w=round(width*20),
            h=1,hRule='atLeast',hAnchor='page',vAnchor='page',x=round(x*20),y=round(y*20),wrap='none'))
    title=marker.insert_paragraph_before('毕业设计（论文）诚信声明')
    position(title,42*72/25.4,131.2,132*72/25.4)
    title.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing=Pt(26)
    set_run(title.runs[0],16,'SimHei',True)
    paragraphs=[
      '本人郑重声明：所呈交的毕业设计（论文）是我个人在导师指导下进行的研究工作及取得的研究成果。就我所知，除了文中特别加以标注和致谢的地方外，论文中不包含其他人已经发表和撰写的研究成果，也不包含为获得华东交通大学或其他教育机构的学位或证书所使用过的材料。',
      '如在文中涉及抄袭或剽窃行为，本人愿承担由此而造成的一切后果及责任。']
    for index,text in enumerate(paragraphs):
        p=marker.insert_paragraph_before(text)
        p.paragraph_format.left_indent=Mm(18);p.paragraph_format.right_indent=Mm(17)
        p.paragraph_format.space_before=Pt(118.134 if index==0 else 0)
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.first_line_indent=Pt(32);p.paragraph_format.line_spacing=Pt(31.2)
        set_run(p.runs[0],16,'KaiTi',True)
    for text,y in [('本人签名',534),('导师签名',596.4)]:
        p=marker.insert_paragraph_before(text)
        position(p,240.9,y,200)
        p.paragraph_format.line_spacing=Pt(31.2)
        set_run(p.runs[0],16)
    p=marker.insert_paragraph_before(meta['sign-date'])
    position(p,230,663,199.5)
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.paragraph_format.line_spacing=Pt(24)
    set_run(p.runs[0],15)
    marker._p.getparent().remove(marker._p)


def field(p,code):
    for tag,typ in [('begin','fldChar'),(code,'instrText'),('separate','fldChar'),('end','fldChar')]:
        r=p.add_run()
        n=OxmlElement('w:'+typ)
        if typ=='fldChar':n.set(qn('w:fldCharType'),tag)
        else:n.text=tag
        r._r.append(n)


def format_document(path,pdf_path,meta):
    doc=Document(path);styles(doc)
    for s in doc.sections:configure_section(s)
    template=deepcopy(doc.sections[-1]._sectPr)
    markers={p.text:p for p in doc.paragraphs if p.text.startswith(MARK)}
    with fitz.open(pdf_path) as pdf:
        add_cover(markers[MARK+'COVER'],pdf)
        figures=[fitz.Rect(r) for page in pdf for r in figure_regions(page, include_images=True)]
    for code in ('DECLARATION','CN','EN','TOC','BODY'):
        break_before(markers[MARK+code],template)
    add_declaration(markers[MARK+'DECLARATION'],doc,meta)
    toc=markers[MARK+'TOC']
    toc.insert_paragraph_before('目　录','TOC Heading')
    for code in ('CN','EN','BODY'):
        p=markers[MARK+code]._p;p.getparent().remove(p)
    for p in list(doc.paragraphs):
        if p.style.name=='Equation':
            width=doc.sections[-1].page_width-doc.sections[-1].left_margin-doc.sections[-1].right_margin
            p.paragraph_format.tab_stops.add_tab_stop(Pt(width/12700/2),WD_TAB_ALIGNMENT.CENTER)
            p.paragraph_format.tab_stops.add_tab_stop(Pt(width/12700),WD_TAB_ALIGNMENT.RIGHT)
            math=p._p.find(qn('m:oMath'))
            if math is None:
                raise ValueError('编号公式未生成原生 Word 公式。')
            leading=OxmlElement('w:r');leading.append(OxmlElement('w:tab'))
            trailing=OxmlElement('w:r');trailing.append(OxmlElement('w:tab'))
            p._p.insert(list(p._p).index(math),leading)
            p._p.insert(list(p._p).index(math)+1,trailing)
        if p.text==MARK+'PAGEBREAK':
            p.clear()
            p.paragraph_format.line_spacing=Pt(1)
            p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(0)
            p.add_run().add_break(WD_BREAK.PAGE)
        if p._p.xpath('.//w:drawing'):
            if not p._p.xpath('.//wp:anchor'):
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent=Pt(0)
                p.paragraph_format.line_spacing=1
                p.paragraph_format.space_before=Pt(12)
                p.paragraph_format.keep_with_next=True
                rect=figures.pop(0)
                for extent in p._p.xpath('.//wp:extent | .//a:xfrm/a:ext'):
                    extent.set('cx',str(round(rect.width*12700)))
                    extent.set('cy',str(round(rect.height*12700)))
    # 真正的表格只有 tabular 内容。
    for table in doc.tables:
        table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        widths=[20,33.5,55.75,55.75] if len(table.columns)==4 else [165/len(table.columns)]*len(table.columns)
        for column,width in zip(table.columns,widths):column.width=Mm(width)
        borders=node('w:tblBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'):
            borders.append(node('w:'+edge,val='single',sz=4,color='808080'))
        pr=table._tbl.tblPr
        old=pr.find(qn('w:tblBorders'))
        if old is not None:pr.remove(old)
        pr.append(borders)
        for row in table.rows:
            row.height=None
            for cell,width in zip(row.cells,widths):
                cell.width=Mm(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tc=cell._tc.get_or_add_tcPr()
                margins=node('w:tcMar')
                for edge in ('top','bottom','left','right'):
                    margins.append(node('w:'+edge,w=60,type='dxa'))
                tc.append(margins)
                for p in cell.paragraphs:
                    p.style=doc.styles['Normal'];p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.first_line_indent=Pt(0)
                    p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(0)
                    p.paragraph_format.line_spacing=Pt(17.0625)
                    for r in p.runs:set_run(r,10.5)
        table.rows[0]._tr.get_or_add_trPr().append(node('w:tblHeader'))
    doc.settings.odd_and_even_pages_header_footer=True
    for i,s in enumerate(doc.sections):
        configure_section(s)
        for hf in (s.header,s.even_page_header,s.footer,s.even_page_footer):
            hf.is_linked_to_previous=False
            hf.paragraphs[0].clear()
        if i!=len(doc.sections)-1:continue
        s._sectPr.append(node('w:pgNumType',start=1))
        for header,text in [(s.header,'华东交通大学毕业设计'),(s.even_page_header,meta['author']+'：'+meta['title'])]:
            p=header.paragraphs[0]
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.line_spacing=Pt(12);p.paragraph_format.space_after=Pt(0)
            set_run(p.add_run(text),9)
            borders=node('w:pBdr');borders.append(node('w:bottom',val='single',sz=3,space=1,color='000000'))
            p._p.get_or_add_pPr().append(borders)
        for footer in (s.footer,s.even_page_footer):
            p=footer.paragraphs[0]
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.line_spacing=Pt(12);p.paragraph_format.space_after=Pt(0)
            field(p,' PAGE ')
            for r in p.runs:set_run(r,9)
    # Word 2013+ 布局，关闭文档网格对标题和正文的附加间距。
    compression=doc.settings.element.find(qn('w:characterSpacingControl'))
    if compression is None:
        compression=node('w:characterSpacingControl',val='compressPunctuation')
        doc.settings.element.append(compression)
    else:
        compression.set(qn('w:val'),'compressPunctuation')
    compat=doc.settings.element.find(qn('w:compat'))
    if compat is not None:
        for n in list(compat):
            if n.get(qn('w:name'))=='compatibilityMode':compat.remove(n)
        compat.append(node('w:compatSetting',name='compatibilityMode',uri='http://schemas.microsoft.com/office/word',val=15))
    doc.core_properties.title=meta['title'];doc.core_properties.author=meta['author']
    doc.core_properties.comments=''
    doc.save(path)


def finalize_word(path,render_path,create_toc=True,refresh_toc=False):
    import win32com.client
    app=win32com.client.DispatchEx('Word.Application')
    app.Visible=False;app.DisplayAlerts=0
    d=None
    try:
        d=app.Documents.Open(str(path),ReadOnly=not (create_toc or refresh_toc),AddToRecentFiles=False)
        if create_toc:
            # Word may open Pandoc's package in its legacy compatibility mode
            # despite the OOXML setting. Set the live document mode as well.
            d.SetCompatibilityMode(15)
            r=d.Content
            if not r.Find.Execute(FindText=MARK+'TOC'):
                raise ValueError('未找到目录插入点。')
            r.Text=''
            toc=d.TablesOfContents.Add(r,UseHeadingStyles=True,UpperHeadingLevel=1,LowerHeadingLevel=3,
                UseFields=False,RightAlignPageNumbers=True,IncludePageNumbers=True,UseHyperlinks=True,
                HidePageNumbersInWeb=False)
            # 使用内置样式 ID，避免中文 Word 把自建的 TOC 1 改名为 TOC 11。
            for level in (1,2,3):
                s=d.Styles(-19-level)
                s.Font.Name='Times New Roman';s.Font.NameFarEast='SimHei' if level==1 else 'SimSun'
                s.Font.Size=12;s.Font.Bold=-1 if level==1 else 0;s.Font.Italic=0;s.Font.Color=0
                p=s.ParagraphFormat
                p.CharacterUnitLeftIndent=0;p.CharacterUnitFirstLineIndent=0
                p.LeftIndent=(level-1)*12;p.FirstLineIndent=0;p.RightIndent=0
                p.SpaceBefore=5 if level==1 else 0;p.SpaceAfter=0
                p.LineSpacingRule=4;p.LineSpacing=19.5;p.Alignment=0
            d.Repaginate();toc.Update();d.Repaginate();toc.UpdatePageNumbers()
            d.Save()
        elif refresh_toc:
            d.Repaginate()
            for toc in d.TablesOfContents:
                toc.UpdatePageNumbers()
            d.Save()
        d.Repaginate()
        d.ExportAsFixedFormat(str(render_path),17,Item=0,DocStructureTags=False)
        print('Native Word pages:',d.ComputeStatistics(2),flush=True)
    finally:
        if d is not None:d.Close(0)
        app.Quit()


def calibrate_chapter_spacing(path, source_pdf, word_pdf):
    """Match chapter baselines and the following paragraph to the PDF.

    Word and TeX distribute the leading above a chapter differently. Keep
    native paragraphs and their line spacing, adjusting glyph position and
    the inter-paragraph gap using the actual rendered baselines.
    """
    from compare_word_pdf import characters
    import unicodedata
    def compact(value):
        return ''.join(unicodedata.normalize('NFKC',value).split())
    def page_data(pdf):
        result=[]
        for page in pdf:
            chars=characters(page)
            result.append((''.join(c[0] for c in chars),chars))
        return result
    doc=Document(path)
    changes=[]
    with fitz.open(source_pdf) as source, fitz.open(word_pdf) as rendered:
        source_pages,word_pages=page_data(source),page_data(rendered)
        targets={compact(title):page-1 for level,title,page in source.get_toc() if level==1}
        for i,p in enumerate(doc.paragraphs):
            if p.style.name!='Heading 1':continue
            title=compact(p.text)
            page=targets.get(title)
            if page is None:continue
            source_text,source_chars=source_pages[page]
            start=source_text.find(title)
            if start<0:continue
            # A TOC occurrence is followed by its page number, whereas the
            # chapter is followed by the next paragraph's actual text.
            following=next((q for q in doc.paragraphs[i+1:] if q.text.strip()),None)
            if following is None:continue
            prefix=compact(following.text)[:24]
            next_start=source_text.find(prefix,start+len(title))
            if not prefix or next_start<0:continue
            candidates=[]
            for actual_page,(text,chars) in enumerate(word_pages):
                at=text.find(title)
                while at>=0:
                    after=text.find(prefix,at+len(title))
                    if after==at+len(title):
                        candidates.append((actual_page,chars[at][1][1],chars[after][1][1]))
                    at=text.find(title,at+1)
            if len(candidates)!=1:continue
            actual_page,title_y,next_y=candidates[0]
            target_title_y=source_chars[start][1][1]
            target_next_y=source_chars[next_start][1][1]
            title_delta=title_y-target_title_y
            gap_delta=next_y-target_next_y
            after=p.paragraph_format.space_after
            if after is None:after=p.style.paragraph_format.space_after
            new_after=(after.pt if after is not None else 0)-gap_delta
            # Large offsets indicate a page/content mismatch, not leading.
            if actual_page!=page or abs(title_delta)>15 or not 0<=new_after<=40:continue
            for run in p.runs:
                pr=run._r.get_or_add_rPr()
                old=pr.find(qn('w:position'))
                if old is not None:pr.remove(old)
                pr.append(node('w:position',val=round(title_delta*2)))
            p.paragraph_format.space_after=Pt(new_after)
            changes.append({'title':p.text,'page':page+1,'raise_pt':round(title_delta*2)/2,
                            'space_after_pt':new_after})
    if changes:doc.save(path)
    return changes


def export_source(source,pdf_path,output):
    source,pdf_path,output=map(lambda p:Path(p).resolve(),(source,pdf_path,output))
    from export_word import protect_existing
    if output.suffix.lower()!='.docx' or output in (source,pdf_path):
        raise ValueError('输出应为独立的 .docx 文件。')
    protect_existing(output)
    if not shutil.which('pandoc'):
        raise ValueError('源码导出需要 Pandoc，请先安装 Pandoc。')
    output.parent.mkdir(parents=True,exist_ok=True)
    (ROOT/'tmp').mkdir(exist_ok=True)
    work=Path(tempfile.mkdtemp(prefix='source-word-',dir=ROOT/'tmp'))
    prepared,meta,inputs,headings,tables=prepare_source(source,pdf_path,work)
    manifest=source_manifest(json.loads(prepared.read_text(encoding='utf-8')))
    for name in ('ecjtu-thesis.cls','references.bib','assets/cover-background.png'):
        path=source.parent/name
        if path.exists():inputs[str(path)]=hashlib.sha256(path.read_bytes()).hexdigest()
    stale=[Path(name).name for name in inputs if Path(name).suffix not in ('.aux','.bbl')
           and Path(name).stat().st_mtime>pdf_path.stat().st_mtime+1]
    if stale:raise ValueError('PDF 早于当前源码，请先构建论文：'+', '.join(stale))
    temp=work/'native.docx'
    subprocess.run(['pandoc',str(prepared),'-f','json','-t','docx',
                    '--syntax-highlighting=none','-o',str(temp)],check=True)
    format_document(temp,pdf_path,meta)
    finalize_word(temp,work/'word.pdf')
    from word_line_layout import restore_line_breaks
    layout_doc=Document(temp)
    line_layout=restore_line_breaks(layout_doc,pdf_path,tracking_twips=-5,word_pdf=work/'word.pdf')
    from word_english_layout import restore_english_hyphens
    english_layout=restore_english_hyphens(layout_doc,pdf_path)
    if english_layout:
        english=[p for p in layout_doc.paragraphs if p.style.name=='English Body']
        english[0].paragraph_format.right_indent=Pt(10)
        english[1].paragraph_format.space_after=Pt(9)
        title=next(p for p in layout_doc.paragraphs if p.style.name=='Front Heading English')
        title.paragraph_format.space_after=Pt(.3)
    # A single line at the bottom of page 16 needs a shorter line box so it
    # remains on the same page as in the source PDF.
    late_line=next((p for p in layout_doc.paragraphs if p.text.startswith(
        '2. 详细设计阶段重点写“打算如何构建”')),None)
    if late_line is not None:
        late_line.paragraph_format.line_spacing=Pt(15)
    if line_layout or english_layout or late_line is not None:
        layout_doc.save(temp)
        finalize_word(temp,work/'word.pdf',create_toc=False,refresh_toc=True)
    chapter_calibration=calibrate_chapter_spacing(temp,pdf_path,work/'word.pdf')
    print('WORK',work,flush=True)
    # 存下源文件关联与实际段落，供后续核验和覆盖保护使用。
    doc=Document(temp)
    from word_toc_layout import format_toc_entries,calibrate_toc_baselines
    toc_calibration=calibrate_toc_baselines(doc,pdf_path,work/'word.pdf')
    format_toc_entries(doc)
    mapping={'schema':3,'mode':'source','pdf':str(pdf_path),
        'pdf_sha256':hashlib.sha256(pdf_path.read_bytes()).hexdigest(),'docx':str(output),
        'source':str(source),'source_sha256':inputs,'headings':headings,'table_count':tables,
        'source_content':manifest,
        'chapter_calibration':chapter_calibration,
        'toc_calibration':toc_calibration,
        'line_layout':line_layout,
        'english_layout':english_layout,
        'paragraphs':[],'pages':[],'work':str(work)}
    for i,p in enumerate(doc.element.xpath('//w:body//w:p[w:r/w:t or .//m:oMath]')):
        if p.xpath('./w:pPr/w:pStyle[starts-with(@w:val,"TOC")]'):continue
        name=f'paragraph_{i}'
        start=node('w:bookmarkStart',id=10000+i,name=name);end=node('w:bookmarkEnd',id=10000+i)
        p.insert(1 if p.pPr is not None else 0,start);p.append(end)
        text=''.join(p.xpath('.//w:t/text() | .//m:t/text()'))
        mapping['paragraphs'].append({'bookmark':name,'text':text,'in_table':bool(p.xpath('ancestor::w:tc'))})
    doc.save(temp)
    # 对最终写入书签的文件重新打开渲染，核验的字节与交付一致。
    finalize_word(temp,work/'word.pdf',create_toc=False,refresh_toc=True)
    from word_baseline_layout import calibrate_paragraph_baselines
    calibrated_doc=Document(temp)
    mapping['paragraph_calibration']=calibrate_paragraph_baselines(calibrated_doc,pdf_path,work/'word.pdf')
    if mapping['paragraph_calibration']:
        calibrated_doc.save(temp)
        finalize_word(temp,work/'word.pdf',create_toc=False,refresh_toc=True)
    mapping['docx_sha256']=hashlib.sha256(temp.read_bytes()).hexdigest()
    with fitz.open(work/'word.pdf') as pdf:
        colored=[(i+1,s['text']) for i,p in enumerate(pdf) for b in p.get_text('dict')['blocks']
                 for line in b.get('lines',[]) for s in line['spans'] if s['text'].strip() and s['color']!=0]
        if colored:raise ValueError('Word 渲染存在非黑色文字：'+str(colored[:3]))
        mapping['pages']=[{'number':i+1,'width':p.rect.width,'height':p.rect.height} for i,p in enumerate(pdf)]
        for i,p in enumerate(pdf):
            p.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False).save(work/f'page-{i+1:02d}.png')
    with ZipFile(temp) as z:
        mapping['media_sha256']={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
    from check_thesis_word import verify_source_word
    if verify_source_word(temp,mapping,write_report=False):
        raise ValueError('源码 Word 的内容或格式检查失败，未替换目标文件。')
    protect_existing(output)
    # 同目录暂存后原子替换，失败时保留上一版。
    with tempfile.NamedTemporaryFile(prefix='.word-',suffix='.docx',dir=output.parent,delete=False) as f:
        staged=Path(f.name)
    try:
        shutil.copy2(temp,staged)
        staged.replace(output)
    finally:
        staged.unlink(missing_ok=True)
    output.with_suffix('.map.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return mapping
