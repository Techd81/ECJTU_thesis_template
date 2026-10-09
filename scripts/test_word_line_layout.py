"""Regression coverage for PDF line boundaries and native paragraph retention."""
from pathlib import Path
import tempfile
import unittest

import pymupdf as fitz
from docx import Document

from word_line_layout import line_boundaries,normalized,restore_line_breaks


class LineLayoutTests(unittest.TestCase):
    def test_superscripts_do_not_create_line_breaks(self):
        chars=[('x',0,100,72),('1',0,95,80),('.',0,100,86),('y',0,119.5,72)]
        self.assertEqual(line_boundaries(chars),{3})

    def test_changed_wrapping_preserves_paragraph_and_text(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder)
            for name,wrapped in [('reference',True),('word',False)]:
                with fitz.open() as pdf:
                    page=pdf.new_page()
                    page.insert_text((72,90),'Heading')
                    if wrapped:
                        page.insert_text((72,120),'First line.')
                        page.insert_text((72,140),'Second line.')
                    else:page.insert_text((72,120),'First line. Second line.')
                    pdf.set_toc([[1,'Heading',1]])
                    pdf.save(root/(name+'.pdf'))
            doc=Document()
            doc.add_paragraph('Heading','Heading 1')
            para=doc.add_paragraph('First line. Second line.')
            original=normalized(para.text)
            records=restore_line_breaks(doc,root/'reference.pdf',word_pdf=root/'word.pdf')
            self.assertEqual(len(doc.paragraphs),2)
            self.assertEqual(normalized(para.text),original)
            self.assertEqual(para.text,'First line.\nSecond line.')
            self.assertEqual(records[0]['lines'],2)
            unchanged=Document()
            unchanged.add_paragraph('Heading','Heading 1')
            unchanged.add_paragraph('First line. Second line.')
            self.assertEqual(restore_line_breaks(unchanged,root/'reference.pdf',word_pdf=root/'reference.pdf'),[])


if __name__=='__main__':unittest.main()
