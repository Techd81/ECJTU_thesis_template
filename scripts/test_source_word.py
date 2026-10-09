"""源码 Word 排版退化回归。集成用例需要已构建 PDF、Pandoc 和 Microsoft Word。"""
from contextlib import redirect_stdout
from copy import deepcopy
import hashlib
from io import StringIO
import json
from pathlib import Path
import shutil
import sys
import tempfile
from types import SimpleNamespace
import unittest

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm,Pt,RGBColor
import pymupdf as fitz

from check_thesis_word import verify_source_word
from export_thesis_word import (export_source,source_manifest,ROOT,pandoc,
                               macro_definitions,expand_listings,
                               normalize_longtables,normalize_table_columns,text_of,
                               bibliography_entries)
from verify_word import visible_text, range_content_text
from compare_word_pdf import characters
from word_toc_layout import format_toc_entries


class ManifestTests(unittest.TestCase):
    def test_multiline_bibliography_author_label_keeps_three_entries(self):
        raw=r'''\begin{thebibliography}{3}
\bibitem[First(2016)]{first} First book.
\bibitem[Second(2017)]{second} Second book.
\bibitem[Third et~al.(2019)Third, Fourth, and
  Fifth]{third} Third book.
\end{thebibliography}'''
        entries=bibliography_entries(raw)
        self.assertEqual([key for key,_ in entries],['first','second','third'])
        self.assertEqual(entries[1][1].strip(),'Second book.')
        self.assertEqual(entries[2][1].strip(),'Third book.')

    def test_unnumbered_toc_entry_uses_page_number_tab(self):
        doc=Document()
        doc.styles.add_style('toc 1',1)
        plain=doc.add_paragraph('致谢\t9','toc 1')
        numbered=doc.add_paragraph('第1章　绪论\t1','toc 1')
        format_toc_entries(doc)
        self.assertEqual(len(plain.paragraph_format.tab_stops),1)
        self.assertGreater(plain.paragraph_format.tab_stops[0].position.pt,460)
        self.assertEqual(len(numbered.paragraph_format.tab_stops),2)
        self.assertIn('\t',numbered.text)

    def test_word_math_serialization_keeps_original_operators(self):
        word_range=SimpleNamespace(OMaths=SimpleNamespace(Count=1), Text='\t𝑥−𝑦\t(5-1)',
            WordOpenXML='''<pkg:package xmlns:pkg="http://schemas.microsoft.com/office/2006/xmlPackage"
                xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
                xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
                <pkg:part><pkg:xmlData><w:document><w:body><w:p>
                <m:oMath><m:r><m:t>x-y</m:t></m:r></m:oMath><w:r><w:t>(5-1)</w:t></w:r>
                </w:p></w:body></w:document></pkg:xmlData></pkg:part></pkg:package>''')
        self.assertEqual(range_content_text(word_range),'x-y(5-1)')
        word_range.WordOpenXML=word_range.WordOpenXML.replace('x-y','x+y')
        self.assertNotEqual(range_content_text(word_range),'x-y(5-1)')
        word_range.OMaths.Count=0
        word_range.Text='普通正文的减号 −'
        self.assertEqual(range_content_text(word_range),word_range.Text)

    def test_code_blocks_are_part_of_the_source_manifest(self):
        ast={'blocks':[{'t':'CodeBlock','c':[['',['python'],[]],'print("original code")']}]}
        self.assertEqual(source_manifest(ast)['code_blocks'],['print("original code")'])

    @unittest.skipUnless(shutil.which('pandoc'),'需要 Pandoc')
    def test_author_macro_values_reach_text_and_math(self):
        definitions=macro_definitions(r'\newcommand{\Cases}{12}\newcommand{\Rate}{4.76}')
        ast=json.loads(pandoc(definitions+r'Tests: \Cases{}. $U=\Rate{}\%$'))
        self.assertIn('Tests: 12.',text_of(ast['blocks']))
        self.assertIn(r'U=4.76\%',text_of(ast['blocks']).replace('{}',''))
        self.assertNotIn('Cases',text_of(ast['blocks']))

    @unittest.skipUnless(shutil.which('pandoc'),'需要 Pandoc')
    def test_custom_table_column_preserves_all_cells(self):
        tex=r'\begin{tabular}{@{}p{24mm}p{48mm}Y@{}} A & B & C\\ D & E & F\\ \end{tabular}'
        ast=json.loads(pandoc(normalize_table_columns(tex)))
        self.assertEqual(len(ast['blocks']),1)
        self.assertEqual(ast['blocks'][0]['t'],'Table')
        self.assertEqual(len(ast['blocks'][0]['c'][2]),3)
        self.assertIn('F',text_of(ast['blocks']))

    @unittest.skipUnless(shutil.which('pandoc'),'需要 Pandoc')
    def test_longtable_keeps_data_once_and_uses_compiled_number(self):
        tex=r'''\begin{longtable}{ll}
\caption{Cases}\label{tab:cases}\\
Id & Result\\ \endfirsthead
Repeated & Header\\ \endhead
\bottomrule \endfoot
T01 & Pass\\ T02 & Pass\\
\end{longtable}'''
        normalized=normalize_longtables(tex,{'tab:cases':'5-1'})
        ast=json.loads(pandoc(normalized))
        text=text_of(ast['blocks'])
        self.assertIn('TABLECAP表5-1',text)
        self.assertNotIn('Repeated',text)
        self.assertEqual(text.count('T01'),1)
        self.assertEqual(text.count('T02'),1)
        self.assertEqual(sum(b['t']=='Table' for b in ast['blocks']),1)

    @unittest.skipUnless(shutil.which('pandoc'),'需要 Pandoc')
    def test_included_code_is_read_and_fingerprinted(self):
        with tempfile.TemporaryDirectory(prefix='code-listing-') as folder:
            root=Path(folder)
            code=root/'core.py'
            code.write_text('print("saved code")\n',encoding='utf-8')
            inputs={}
            expanded=expand_listings(r'\lstinputlisting[language=Python]{core.py}',root,inputs)
            manifest=source_manifest(json.loads(pandoc(expanded)))
            self.assertEqual(manifest['code_blocks'],['print("saved code")'])
            self.assertEqual(inputs[str(code.resolve())],hashlib.sha256(code.read_bytes()).hexdigest())

    def test_quoted_words_and_math_stay_separate(self):
        ast={'blocks':[{'t':'Para','c':[{'t':'Quoted','c':[{'t':'DoubleQuote'},[{'t':'Str','c':'原文'}]]},
            {'t':'Math','c':[{'t':'InlineMath'},r'\to']},{'t':'Str','c':'结果'}]}]}
        self.assertEqual(source_manifest(ast),{'paragraphs':['“原文”结果'],'math_count':1})

    def test_fake_bold_is_visible_but_hidden_text_is_not(self):
        with fitz.open() as pdf:
            p=pdf.new_page()
            for mode in (0,1):p.insert_text((72,72),'ABC',render_mode=mode)
            p.insert_text((95,72),'123')
            p.insert_text((72,100),'invisible',render_mode=3)
            self.assertEqual(visible_text(pdf),'ABC123')


@unittest.skipUnless(sys.platform=='win32' and shutil.which('pandoc') and (ROOT/'build/main.pdf').exists(),
                     '需要 Windows、Pandoc 和已编译的论文 PDF')
class SourceWordTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory=tempfile.TemporaryDirectory(prefix='source-word-test-')
        cls.root=Path(cls.directory.name)
        cls.original=cls.root/'original.docx'
        with redirect_stdout(StringIO()):
            cls.mapping=export_source(ROOT/'main.tex',ROOT/'build/main.pdf',cls.original)

    @classmethod
    def tearDownClass(cls):cls.directory.cleanup()

    def setUp(self):
        self.path=self.root/'mutated.docx'
        shutil.copy2(self.original,self.path)
        self.mapping=deepcopy(type(self).mapping)

    def check(self,doc=None):
        if doc is not None:doc.save(self.path)
        self.mapping['docx_sha256']=hashlib.sha256(self.path.read_bytes()).hexdigest()
        with redirect_stdout(StringIO()):result=verify_source_word(self.path,self.mapping)
        report=json.loads(self.path.with_suffix('.check.json').read_text(encoding='utf-8'))
        return result,{c['name'] for c in report['checks'] if not c['passed']}

    def test_native_template_passes(self):
        self.assertEqual(self.check()[0],0)

    def test_matching_chapter_baselines_match_pdf(self):
        matching_pages=0
        with fitz.open(self.mapping['pdf']) as source, fitz.open(Path(self.mapping['work'])/'word.pdf') as word:
            for level,title,target in source.get_toc():
                if level!=1 or target>len(word):continue
                page=target-1
                a,b=characters(source[page]),characters(word[page])
                source_text=''.join(c[0] for c in a)
                word_text=''.join(c[0] for c in b)
                title=''.join(title.split())
                left,right=source_text.find(title),word_text.find(title)
                if left<0 or right<0:continue
                matching_pages+=1
                with self.subTest(title=title):
                    self.assertLess(abs(a[left][1][1]-b[right][1][1]),.3)
        self.assertGreaterEqual(matching_pages,5)

    def test_legacy_layout_mode_is_rejected(self):
        d=Document(self.path)
        d.settings.element.xpath('./w:compat/w:compatSetting[@w:name="compatibilityMode"]')[0].set(qn('w:val'),'12')
        self.assertIn('Word 2013 排版兼容模式',self.check(d)[1])

    def test_disabled_punctuation_compression_is_rejected(self):
        d=Document(self.path)
        d.settings.element.xpath('./w:characterSpacingControl')[0].set(qn('w:val'),'doNotCompress')
        self.assertIn('中文标点压缩',self.check(d)[1])

    def test_wrong_body_margins_are_rejected(self):
        d=Document(self.path);d.sections[-1].top_margin=Mm(8)
        self.assertIn('第 6 节 A4 与页边距',self.check(d)[1])

    def test_theme_override_is_rejected(self):
        d=Document(self.path)
        d.styles['Heading 1'].element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsiaTheme'),'majorEastAsia')
        self.assertIn('Heading 1 字体不受主题覆盖',self.check(d)[1])

    def test_caption_italics_are_rejected(self):
        d=Document(self.path);d.styles['Image Caption'].font.italic=True
        self.assertIn('Image Caption 直立字体',self.check(d)[1])

    def test_merged_paragraphs_are_rejected_by_source(self):
        d=Document(self.path)
        p=next(p for p in d.paragraphs if p.style.name=='Abstract Body')
        following=p._p.getnext()
        for r in list(following):
            if r.tag!=qn('w:pPr'):p._p.append(r)
        following.getparent().remove(following)
        self.assertIn('全部源段落文字与边界',self.check(d)[1])

    def test_extra_heading_number_is_rejected(self):
        d=Document(self.path)
        p=next(p for p in d.paragraphs if p.style.name=='Heading 3')
        p.runs[0].text='99. '+p.runs[0].text
        self.assertIn('与 LaTeX 目录独立对照',self.check(d)[1])

    def test_shrunken_diagram_is_rejected(self):
        d=Document(self.path)
        d.element.xpath('//wp:inline/wp:extent')[0].set('cx','5334000')
        self.assertIn('流程图 1 物理尺寸',self.check(d)[1])

    def test_changed_source_record_is_rejected(self):
        self.mapping['source_sha256'][str(ROOT/'main.tex')]='bad'
        self.assertIn('源文件 main.tex',self.check()[1])

    def test_blue_hyperlink_is_rejected(self):
        d=Document(self.path);d.styles['Hyperlink'].font.color.rgb=RGBColor(0,0,255)
        self.assertIn('Hyperlink 链接黑色',self.check(d)[1])

    def test_missing_reference_hanging_indent_is_rejected(self):
        d=Document(self.path);d.styles['Bibliography'].paragraph_format.first_line_indent=Pt(0)
        self.assertIn('Bibliography 字号字体行距缩进',self.check(d)[1])

    def test_missing_compiled_reference_is_rejected(self):
        # 模拟源 AST 静默漏掉编译所得的一条文献，须由 BBL 独立检查发现。
        bbl=self.root/'additional-reference.bbl'
        bbl.write_text(r'\bibitem[Author(2026)]{missing} Omitted reference.',encoding='utf-8')
        self.mapping['source_sha256'][str(bbl)]=hashlib.sha256(bbl.read_bytes()).hexdigest()
        self.assertIn('编译参考文献条目与编号',self.check()[1])


if __name__=='__main__':unittest.main()
