"""Word 正文可见性、内容完整性及审阅记录保护回归测试。"""
from contextlib import redirect_stdout
import hashlib
from io import StringIO, BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from lxml import etree
from PIL import Image
import pymupdf as fitz

from check_word import verify, NS
from export_word import export


class WordExportTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix='word-export-test-')
        self.work = Path(self.directory.name)
        self.pdf = self.work/'paper.pdf'
        self.docx = self.work/'review.docx'
        with fitz.open() as pdf:
            for text in ('First page: scope and purpose.', 'Second page: results and references.'):
                page = pdf.new_page(width=595.28, height=841.89)
                page.insert_text((72, 90), text, fontsize=12)
            image = BytesIO()
            Image.new('RGB', (16, 16), 'blue').save(image, format='PNG')
            pdf[0].insert_image(fitz.Rect(72, 120, 104, 152), stream=image.getvalue())
            pdf.save(self.pdf)
        with redirect_stdout(StringIO()):
            export(self.pdf, self.docx, dpi=150)

    def tearDown(self):
        self.directory.cleanup()

    def check(self):
        with redirect_stdout(StringIO()):
            return verify(self.docx)

    def modify_xml(self, change):
        with ZipFile(self.docx) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        root = etree.fromstring(parts['word/document.xml'])
        change(root)
        parts['word/document.xml'] = etree.tostring(root)
        with ZipFile(self.docx, 'w', ZIP_DEFLATED) as archive:
            for name, data in parts.items():
                archive.writestr(name, data)
        # 更新容器哈希，确保内容检查本身能检出损坏，而非仅依赖哈希。
        map_path = self.docx.with_suffix('.map.json')
        mapping = json.loads(map_path.read_text(encoding='utf-8'))
        mapping['docx_sha256'] = hashlib.sha256(self.docx.read_bytes()).hexdigest()
        map_path.write_text(json.dumps(mapping), encoding='utf-8')

    def test_two_pages_and_searchable_text(self):
        self.assertEqual(self.check(), 0)
        text = '\n'.join(p.text for p in Document(self.docx).paragraphs)
        self.assertIn('First page: scope and purpose.', text)
        self.assertIn('Second page: results and references.', text)

    def test_missing_character_is_detected(self):
        def remove(root):
            text = root.xpath('//w:t', namespaces=NS)[0]
            text.text = ''
        self.modify_xml(remove)
        self.assertEqual(self.check(), 1)

    def test_transparent_text_is_detected(self):
        def hide(root):
            props = root.xpath('//w:r[w:t]/w:rPr', namespaces=NS)[0]
            fill = etree.SubElement(props, '{'+NS['w14']+'}textFill')
            solid = etree.SubElement(fill, '{'+NS['w14']+'}solidFill')
            color = etree.SubElement(solid, '{'+NS['w14']+'}srgbClr')
            alpha = etree.SubElement(color, '{'+NS['w14']+'}alpha')
            alpha.set('{'+NS['w14']+'}val', '100000')
        self.modify_xml(hide)
        self.assertEqual(self.check(), 1)

    def test_missing_image_is_detected(self):
        def remove(root):
            blip = root.xpath('//a:blip', namespaces=NS)[0]
            blip.set('{'+NS['r']+'}embed', 'missingImage')
        self.modify_xml(remove)
        self.assertEqual(self.check(), 1)

    def test_changed_pdf_is_detected(self):
        with fitz.open(self.pdf) as pdf:
            pdf.set_metadata({'title': 'changed source'})
            pdf.saveIncr()
        self.assertEqual(self.check(), 1)

    def test_existing_teacher_comments_are_preserved(self):
        doc = Document(self.docx)
        paragraph = next(p for p in doc.paragraphs if p.text)
        doc.add_comment(paragraph.runs[0], text='Please explain this.', author='Reviewer')
        doc.save(self.docx)
        before = self.docx.read_bytes()
        with self.assertRaisesRegex(ValueError, '已有批注'):
            export(self.pdf, self.docx)
        self.assertEqual(self.docx.read_bytes(), before)

    def test_edited_body_is_preserved(self):
        doc = Document(self.docx)
        doc.add_paragraph('Teacher changed the body.')
        doc.save(self.docx)
        before = self.docx.read_bytes()
        with self.assertRaisesRegex(ValueError, '已修改'):
            export(self.pdf, self.docx)
        self.assertEqual(self.docx.read_bytes(), before)

    def test_white_text_is_detected(self):
        def hide(root):
            props = root.xpath('//w:r[w:t]/w:rPr', namespaces=NS)[0]
            color = etree.SubElement(props, '{'+NS['w']+'}color')
            color.set('{'+NS['w']+'}val', 'FFFFFF')
        self.modify_xml(hide)
        self.assertEqual(self.check(), 1)

    def test_source_spaces_between_font_spans_survive(self):
        pdf_path = self.work/'spaces.pdf'
        output = self.work/'spaces.docx'
        with fitz.open() as pdf:
            page = pdf.new_page()
            page.insert_text((72, 90), 'alpha', fontsize=12, fontname='hebo')
            x = 72 + fitz.get_text_length('alpha', fontsize=12, fontname='hebo')
            page.insert_text((x, 90), ' ', fontsize=12, fontname='helv')
            x += fitz.get_text_length(' ', fontsize=12, fontname='helv')
            page.insert_text((x, 90), 'beta', fontsize=12, fontname='hebo')
            pdf.save(pdf_path)
        with redirect_stdout(StringIO()):
            export(pdf_path, output)
        self.assertIn('alpha beta', ''.join(Document(output).element.xpath('//w:t/text()')))

    def test_output_has_native_visible_text(self):
        doc = Document(self.docx)
        self.assertEqual(doc.styles['Normal'].font.size.pt, 12)
        self.assertFalse(doc.element.xpath('//w:framePr | //w14:textFill | //w:vanish'))
        self.assertFalse(doc.element.xpath('//w:trHeight[@w:hRule="exact"]'))

    def test_stretched_background_keeps_pdf_dimensions(self):
        pdf_path = self.work/'cover.pdf'
        output = self.work/'cover.docx'
        image = BytesIO()
        Image.new('RGB', (400, 500), 'lightblue').save(image, format='PNG')
        with fitz.open() as pdf:
            page = pdf.new_page(width=595.28, height=841.89)
            page.insert_image(page.rect, stream=image.getvalue(), keep_proportion=False)
            page.insert_text((72, 400), 'Editable cover title', fontsize=12)
            pdf.save(pdf_path)
        with redirect_stdout(StringIO()):
            export(pdf_path, output)
        extent = Document(output).element.xpath('//wp:anchor/wp:extent')[0]
        self.assertAlmostEqual(int(extent.get('cy')) / 12700, 841.89, places=2)
        self.assertAlmostEqual(int(extent.get('cx')) / 12700, 595.28, places=2)

    def test_unselectable_page_does_not_replace_existing_output(self):
        blank = self.work/'blank.pdf'
        with fitz.open() as pdf:
            pdf.new_page()
            pdf.save(blank)
        before = self.docx.read_bytes()
        with self.assertRaisesRegex(ValueError, '没有可提取的文字'):
            export(blank, self.docx)
        self.assertEqual(self.docx.read_bytes(), before)


if __name__ == '__main__':
    unittest.main()
